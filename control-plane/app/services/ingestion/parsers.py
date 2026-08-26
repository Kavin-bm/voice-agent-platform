from io import BytesIO

import openpyxl
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader


def parse_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def parse_docx(data: bytes) -> str:
    doc = DocxDocument(BytesIO(data))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def parse_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def parse_html(data: bytes) -> str:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "xlsx": parse_xlsx,
    "txt": parse_txt,
    "url": parse_html,
}


def parse(source_type: str, data: bytes) -> str:
    parser = PARSERS.get(source_type)
    if parser is None:
        raise ValueError(f"Unsupported document source_type: {source_type!r}")
    return parser(data)
